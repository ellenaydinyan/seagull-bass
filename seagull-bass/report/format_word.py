"""Small formatting pass on the pandoc Word output: fonts, table borders, image size."""
import sys
import docx
from docx.shared import Pt, Cm, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.enum.text import WD_ALIGN_PARAGRAPH

path = sys.argv[1]
d = docx.Document(path)
styles = {s.style_id: s for s in d.styles}
for sid in ['Normal', 'BodyText', 'FirstParagraph', 'Compact', 'Title', 'Subtitle',
            'Author', 'Date', 'Heading1', 'Bibliography', 'ImageCaption']:
    s = styles.get(sid)
    if s is None:
        continue
    rpr = s.element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts'); rpr.append(rf)
    for a in ['w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs']:
        rf.set(qn(a), 'Calibri')
    for a in ['w:asciiTheme', 'w:hAnsiTheme', 'w:eastAsiaTheme', 'w:cstheme']:
        rf.attrib.pop(qn(a), None)
for sid in ['Normal', 'BodyText', 'FirstParagraph', 'Compact']:
    if sid in styles:
        styles[sid].font.size = Pt(11)
h = styles['Heading1']
h.font.size = Pt(14); h.font.bold = True; h.font.color.rgb = RGBColor(0x1B, 0x7F, 0x5C)
for sid in ['Title', 'Subtitle', 'Author', 'Date']:
    if sid in styles:
        styles[sid].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
for sec in d.sections:
    sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Cm(2.4)
for t in d.tables:
    borders = OxmlElement('w:tblBorders')
    for edge in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        e = OxmlElement(f'w:{edge}')
        e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), '4'); e.set(qn('w:color'), 'A0A0A0')
        borders.append(e)
    t._tbl.tblPr.append(borders)
    for i, row in enumerate(t.rows):
        for c in row.cells:
            for p in c.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(10)
                    if i == 0:
                        r.font.bold = True
            if i == 0:
                shd = OxmlElement('w:shd'); shd.set(qn('w:val'), 'clear'); shd.set(qn('w:fill'), 'E3F2EB')
                c._tc.get_or_add_tcPr().append(shd)
for shp in d.inline_shapes:
    ratio = shp.height / shp.width
    shp.width = Cm(15); shp.height = int(Cm(15) * ratio)
d.save(path)
print(f"{len(d.tables)} tables, {len(d.inline_shapes)} images")
